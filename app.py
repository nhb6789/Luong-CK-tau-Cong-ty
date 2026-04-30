import io
import zipfile

import openpyxl
import streamlit as st
from openpyxl.worksheet.datavalidation import DataValidation


st.set_page_config(page_title="Tach luong ra 5 file", layout="wide")
st.title("Tach du lieu LUONGSOURCE thanh 1.xlsx - 5.xlsx")


SOURCE_DEFAULT = "LUONGSOURCE.xlsx"
TEMPLATE_DEFAULTS = ["1.xlsx", "2.xlsx", "3.xlsx", "4.xlsx", "5.xlsx"]
PAYMENT_SHEET = "Thanh toan luong ngoai he thong"
BANK_SHEET = "Chi nhanh ngan hang huong"
SOURCE_MAIN_SHEET = "CK"


def normalize_text(value):
    if value is None:
        return ""
    return str(value).strip()


def get_source_account_from_ck(ws_ck):
    for r in range(1, min(ws_ck.max_row, 30) + 1):
        text = normalize_text(ws_ck.cell(r, 1).value)
        if "Số tài khoản" in text or "So tai khoan" in text:
            digits = "".join(ch for ch in text if ch.isdigit())
            if digits:
                return digits
    return "04201015059021"


def build_bank_maps(ws_bank):
    code_to_name = {}
    alias_to_name = {}

    for r in range(2, ws_bank.max_row + 1):
        # Mapping tu cot L -> M (du lieu da tinh toan khi doc data_only=True)
        short_code = normalize_text(ws_bank.cell(r, 12).value)
        full_name_from_code = normalize_text(ws_bank.cell(r, 13).value)
        if short_code and full_name_from_code and full_name_from_code != "#N/A":
            code_to_name[short_code.upper()] = full_name_from_code

        # Mapping tu cot J -> B (key 6 ky tu -> ten chi nhanh)
        alias_key = normalize_text(ws_bank.cell(r, 10).value)
        branch_name = normalize_text(ws_bank.cell(r, 2).value)
        if alias_key and branch_name:
            alias_to_name[alias_key.upper()] = branch_name

    return code_to_name, alias_to_name


def parse_sections_from_ck(ws_ck):
    header_rows = []
    for r in range(1, ws_ck.max_row + 1):
        value = ws_ck.cell(r, 1).value
        if isinstance(value, str) and normalize_text(value).upper() == "STT":
            header_rows.append(r)

    sections = []
    for i, start_header in enumerate(header_rows):
        next_header = header_rows[i + 1] if i + 1 < len(header_rows) else ws_ck.max_row + 1
        data_rows = []
        r = start_header + 1
        while r < next_header:
            stt_value = ws_ck.cell(r, 1).value
            if stt_value is None:
                r += 1
                continue
            # Du lieu hop le la dong co STT dang so
            if isinstance(stt_value, (int, float)) or (isinstance(stt_value, str) and stt_value.strip().isdigit()):
                data_rows.append(r)
            r += 1

        sections.append(
            {
                "header_row": start_header,
                "data_rows": data_rows,
            }
        )
    return sections


def to_account_string(value):
    if value is None:
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        # Tai khoan ngan hang khong can phan thap phan
        if value.is_integer():
            return str(int(value))
    return str(value).strip()


def resolve_bank_name(bank_text, code_hint, code_to_name, alias_to_name):
    code_hint_norm = normalize_text(code_hint).upper()
    if code_hint_norm in code_to_name:
        return code_to_name[code_hint_norm]

    bank_text_norm = normalize_text(bank_text)
    if not bank_text_norm:
        return ""

    # Thu map theo 6 ky tu dau (giong cot J cua bang ma ngan hang)
    alias_key = bank_text_norm.upper()[:6]
    if alias_key in alias_to_name:
        return alias_to_name[alias_key]

    # Fallback theo tu khoa pho bien de format giong file mau
    bank_upper = bank_text_norm.upper()
    fallback_rules = [
        ("VIETCOM", "VIETCOMBANK VIET NAM"),
        ("VIETIN", "VIETINBANK VIET NAM"),
        ("TECHCOM", "TECHCOMBANK HOI SO CHINH"),
        ("AGRIBANK", "AGRIBANK HOI SO CHINH"),
        ("COOP", "NH HOP TAC (CO-OPBANK)"),
        ("BIDV", "BIDV HOAN KIEM"),
        ("MB", "MB HA NOI (HN)"),
        ("ACB", "ACB HA NOI (HN)"),
    ]
    for keyword, bank_name in fallback_rules:
        if keyword in bank_upper:
            return bank_name

    # Neu van khong map duoc thi lay phan truoc dau phay
    return bank_text_norm.split(",")[0].strip().upper()


def clear_payment_rows(ws_payment, start_row=2, end_col=7):
    for r in range(start_row, ws_payment.max_row + 1):
        for c in range(1, end_col + 1):
            ws_payment.cell(r, c).value = None


def fill_payment_sheet(ws_payment, ws_ck, data_rows, source_account, code_to_name, alias_to_name):
    clear_payment_rows(ws_payment)

    written_rows = 0
    for src_row in data_rows:
        beneficiary_name = ws_ck.cell(src_row, 2).value
        beneficiary_account = ws_ck.cell(src_row, 3).value
        bank_text = ws_ck.cell(src_row, 4).value
        amount = ws_ck.cell(src_row, 5).value
        narrative = ws_ck.cell(src_row, 6).value
        bank_code_hint = ws_ck.cell(src_row, 10).value
        resolved_bank = resolve_bank_name(bank_text, bank_code_hint, code_to_name, alias_to_name)

        # Bo qua giao dich noi bo MSB de dung mau "ngoai he thong"
        resolved_bank_upper = normalize_text(resolved_bank).upper()
        if not resolved_bank_upper or resolved_bank_upper == "#N/A" or "MSB" in resolved_bank_upper:
            continue

        written_rows += 1
        ws_payment.cell(written_rows + 1, 1).value = written_rows
        ws_payment.cell(written_rows + 1, 2).value = source_account
        ws_payment.cell(written_rows + 1, 3).value = to_account_string(beneficiary_account)
        ws_payment.cell(written_rows + 1, 4).value = beneficiary_name
        ws_payment.cell(written_rows + 1, 5).value = resolved_bank
        ws_payment.cell(written_rows + 1, 6).value = amount
        ws_payment.cell(written_rows + 1, 7).value = narrative

    return written_rows


def apply_branch_dropdown(ws_payment, ws_bank_values):
    last_row = 2
    for r in range(2, ws_bank_values.max_row + 1):
        if normalize_text(ws_bank_values.cell(r, 2).value):
            last_row = r

    formula = f"='{BANK_SHEET}'!$B$2:$B${last_row}"
    dv = DataValidation(
        type="list",
        formula1=formula,
        allow_blank=True,
        showDropDown=False,
    )
    dv.promptTitle = "Chon ten chi nhanh"
    dv.prompt = "Chon gia tri tu danh sach."
    dv.errorTitle = "Gia tri khong hop le"
    dv.error = "Ten chi nhanh phai nam trong danh sach sheet Chi nhanh ngan hang huong."

    ws_payment.add_data_validation(dv)
    dv.add(f"E2:E{ws_payment.max_row}")


def wb_to_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def build_zip_bytes(files):
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        for file_name, file_bytes in files:
            zf.writestr(file_name, file_bytes)
    zip_buffer.seek(0)
    return zip_buffer.getvalue()


def load_template_pair(template_stream_or_path):
    if isinstance(template_stream_or_path, str):
        wb_write = openpyxl.load_workbook(template_stream_or_path, data_only=False)
        wb_values = openpyxl.load_workbook(template_stream_or_path, data_only=True)
        return wb_write, wb_values

    raw_bytes = template_stream_or_path.getvalue()
    wb_write = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=False)
    wb_values = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
    return wb_write, wb_values


st.markdown(
    "Upload file nguon dang `LUONGSOURCE.xlsx`, app se tach du lieu thanh 5 file theo mau `1.xlsx` den `5.xlsx`."
)

source_file = st.file_uploader(
    "File nguon LUONGSOURCE",
    type=["xlsx"],
    help=f"Neu bo trong se dung file mac dinh: {SOURCE_DEFAULT}",
)

st.subheader("Mau 5 file dich")
tpl_cols = st.columns(5)
template_uploads = []
for i in range(5):
    with tpl_cols[i]:
        template_uploads.append(
            st.file_uploader(
                f"Mau {i + 1}.xlsx",
                type=["xlsx"],
                key=f"template_{i+1}",
                help=f"Mac dinh: {TEMPLATE_DEFAULTS[i]}",
            )
        )

if st.button("Tach du lieu", type="primary"):
    try:
        source_wb = openpyxl.load_workbook(source_file if source_file else SOURCE_DEFAULT, data_only=True)
    except Exception as err:
        st.error(f"Khong mo duoc file nguon: {err}")
        st.stop()

    if SOURCE_MAIN_SHEET not in source_wb.sheetnames:
        st.error(f"File nguon khong co sheet '{SOURCE_MAIN_SHEET}'.")
        st.stop()

    ws_ck = source_wb[SOURCE_MAIN_SHEET]
    sections = parse_sections_from_ck(ws_ck)

    if len(sections) < 5:
        st.error(f"Chi tim thay {len(sections)} nhom du lieu trong sheet CK, can it nhat 5.")
        st.stop()

    source_account = get_source_account_from_ck(ws_ck)
    results = []
    output_files = []

    for idx in range(5):
        template_name = TEMPLATE_DEFAULTS[idx]
        upload_obj = template_uploads[idx]
        template_stream = upload_obj if upload_obj else template_name

        try:
            out_wb, map_wb = load_template_pair(template_stream)
        except Exception as err:
            st.error(f"Khong mo duoc template {template_name}: {err}")
            st.stop()

        if PAYMENT_SHEET not in out_wb.sheetnames or BANK_SHEET not in out_wb.sheetnames:
            st.error(f"Template {template_name} phai co 2 sheet: '{PAYMENT_SHEET}' va '{BANK_SHEET}'.")
            st.stop()

        ws_payment = out_wb[PAYMENT_SHEET]
        ws_bank = map_wb[BANK_SHEET]
        code_to_name, alias_to_name = build_bank_maps(ws_bank)

        filled_rows = fill_payment_sheet(
            ws_payment=ws_payment,
            ws_ck=ws_ck,
            data_rows=sections[idx]["data_rows"],
            source_account=source_account,
            code_to_name=code_to_name,
            alias_to_name=alias_to_name,
        )
        apply_branch_dropdown(ws_payment, ws_bank)

        out_name = f"{idx + 1}_UPDATED.xlsx"
        output_files.append((out_name, wb_to_bytes(out_wb)))
        results.append({"file": out_name, "rows_filled": filled_rows})

    st.success("Da tach va cap nhat xong 5 file.")
    st.write(results)

    zip_bytes = build_zip_bytes(output_files)
    st.download_button(
        label="Tai file ZIP (gom 5 file)",
        data=zip_bytes,
        file_name="LUONG_SPLIT_FILES.zip",
        mime="application/zip",
    )

    for fname, fbytes in output_files:
        st.download_button(
            label=f"Tai {fname}",
            data=fbytes,
            file_name=fname,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
